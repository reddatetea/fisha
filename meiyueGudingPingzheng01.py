'''
每月月底固定凭证，根据工次表生成
2025-7-16加电商工资调整
'''
import re
import os
import easygui
import pandas as pd
import openpyxl

sheet_name = r'工资'
def getDf(msg):
    fname = easygui.fileopenbox(msg=f'请点选{msg}工资表.xlsx')
    df = pd.read_excel(fname, sheet_name=sheet_name)
    return fname

def getPivotTotal(df_,bianma,kemu):
    pivot = pd.pivot_table(df_,index = ['记账科目编码','记账科目','平台'],values = '实发数',aggfunc = 'sum')
    pivot = pivot.reset_index()
    pivot['实发数'] = pivot['实发数'].round(2)
    total = pivot.sum()
    total[0] = bianma
    total[1] = kemu
    total[2] = ''
    total[-1] = pivot['实发数'].sum()*-1
    pivot.loc[bianma] = total
    return pivot

#生产工资
fname_shengcan = getDf('请点选当月“生产人员工资.xlsx”')
# fname_shengcan = r"F:\a00nutstore\008\zw08\gongzi\生产人员工资.xlsx"
path,filename = os.path.split(fname_shengcan)
os.chdir(path)
df_shengcan = pd.read_excel(fname_shengcan,sheet_name = sheet_name)
yingfa = df_shengcan['实发数'].sum() + df_shengcan['代扣社保'].sum() + df_shengcan['扣税'].sum() \
- df_shengcan['补扣税'].sum()

#双佳行管工资
fname = getDf('请点选行管工资')
# fname = r"F:\a00nutstore\008\zw08\gongzi\行管工资.xlsx"
df = pd.read_excel(fname,sheet_name = sheet_name)
gs = '双佳'
df =  df[df['公司'] == gs]
gp = df.groupby('部门')
data = {}
for k,v in gp:
    shebao = gp['社保'].sum()[k]
    yingfashu = gp['应发数'].sum()[k]
    data[k] = shebao+yingfashu
renshu_yanfa = gp['姓名'].get_group('设计研发部').count()
yanglao = round(renshu_yanfa * 588,2)      #养老588元/人
shiye = round(renshu_yanfa * 25.73,2)        #失业25.73元/人
gongshang = round(renshu_yanfa * 60.64,2)        #工伤60.64元/人
yiliao = round(renshu_yanfa * 316.38,2)        #医疗316.38元/人
expense_zhizhao = data['仓库搬运'] + data['生产部']
expense_xingguan = data['行政部'] + data['财务部']
expense_design = data['设计研发部']
expense_sale = data['营销部']

#每月固定凭证
fname_fix = r"F:\a00nutstore\008\zww08\2024\每月固定凭证.xlsx"
wb = openpyxl.load_workbook(fname_fix)
ws = wb['pingzheng']
#写入数据
ws['D5'].value = yingfa
ws['E6'].value = yingfa
ws['D7'].value = expense_zhizhao
ws['D8'].value = expense_xingguan
ws['D9'].value = expense_design
ws['D10'].value = expense_sale
ws['E11'].value = expense_zhizhao + expense_xingguan + expense_design + expense_sale
#社保调整写入
ws['D24'].value = yanglao * -1
ws['D25'].value = yanglao
ws['D26'].value = shiye * -1
ws['D27'].value = shiye
ws['D28'].value = gongshang * -1
ws['D29'].value = gongshang
ws['D30'].value = yiliao * -1
ws['D31'].value = yiliao



#电商工资调整写入
qijian = easygui.enterbox('请输入期间"2025-06"')
newfname = os.path.join(path,f'每月固定凭证{qijian}.xlsx')
wb.save(newfname)
df_new = pd.read_excel(newfname,sheet_name = 'pingzheng')
max_row = df_new.shape[0]


#电商工资分布
# fname_dianshang = easygui.fileopenbox('请点选"电商人员工资分布"')
fname_dianshang = r"F:\a00nutstore\008\zww08\gongzi\电商人员工资分布.xlsx"
df_dianshang_yunying = pd.read_excel(fname_dianshang,sheet_name = '销售费用')
df_dianshang_changku = pd.read_excel(fname_dianshang,sheet_name = '制造费用')
resutl = pd.concat([df_dianshang_yunying,df_dianshang_changku])
#电商工资,工资表简化
df1 = df[['姓名','实发数']]
#运营人员工资
yunying_gongzi = pd.merge(df_dianshang_yunying,df1,how = 'inner',left_on = '姓名',right_on = '姓名')
#仓库人员工资
changku_gongzi = pd.merge(df_dianshang_changku,df1,how = 'inner',left_on = '姓名',right_on = '姓名')
result = pd.concat([yunying_gongzi,changku_gongzi])
result = result.sort_values(by = '记账科目编码')
pivot_yunying = getPivotTotal(yunying_gongzi,'601001','销售费用')
pivot_changku = getPivotTotal(changku_gongzi,'5201003','制造费用')

max_row1 = pivot_yunying.shape[0]

with pd.ExcelWriter(newfname,engine='openpyxl', mode='a', if_sheet_exists='overlay') as writer:
    pivot_yunying.to_excel(writer, sheet_name ='pingzheng', startrow=max_row + 2, index=False)
    pivot_changku.to_excel(writer, sheet_name ='pingzheng', startrow=max_row + 2 + max_row1 + 2, index=False)
    result.to_excel(writer,sheet_name = '明细',index = False)

os.startfile(newfname)





